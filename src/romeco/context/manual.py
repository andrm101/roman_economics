"""Loaders for researcher-supplied datasets (route: manual in config/datasets.yaml).

* Itiner-e roads (CC BY 4.0) -> derived tables may be committed.
* Hanson 2016 Cities Database (licence not stated) -> only aggregates committed;
  row-level extracts stay in data/raw/context/hanson_cities/derived/ (git-ignored).
* Harper 2016 prices/wages/rents (licence not stated) -> same policy as Hanson.
"""

from __future__ import annotations

import hashlib
import math
import sqlite3
import struct
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[3]
RAW = ROOT / "data" / "raw" / "context"
MANUAL = {
    "itinere": RAW / "itinere" / "manual" / "itinere_roads.gpkg",
    "hanson_cities": RAW / "hanson_cities" / "manual" / "hanson2016_citiesdatabase_oxrep.xlsx",
    "harper_prices": RAW / "harper_prices" / "manual" / "Harper 2016 Prices Wages and Rents for Roman Egypt, 1-700.xlsx",
}
EXPECTED_SHA256 = {
    "itinere": "1b9470aaec696b0827e7aecad20badc74df2652ae8a62464baea4f07ae4855e7",
    "hanson_cities": "5b9fc0195c3b5f2df71d6380174974215280cafd3bf5fa463fbf96969b549b6c",
    "harper_prices": "167a62be60e3e28bf116860cf7531b05f597e3ce2c35cf3e04a469ca7fefc34c",
}


def available(name: str) -> bool:
    return MANUAL[name].exists()


def verify(name: str) -> str:
    h = hashlib.sha256(MANUAL[name].read_bytes()).hexdigest()
    if h != EXPECTED_SHA256[name]:
        raise RuntimeError(f"{name}: checksum {h} differs from registered {EXPECTED_SHA256[name]}")
    return h


# ---------------------------------------------------------------------------
# GeoPackage geometry decoding (no GIS dependency)
# ---------------------------------------------------------------------------
_A, _E = 6378137.0, 0.0818191908426215  # WGS84; EPSG:3395 World Mercator (ellipsoidal)


def mercator3395_to_lonlat(x: float, y: float) -> tuple[float, float]:
    lon = math.degrees(x / _A)
    t = math.exp(-y / _A)
    phi = math.pi / 2 - 2 * math.atan(t)
    for _ in range(15):
        es = _E * math.sin(phi)
        nxt = math.pi / 2 - 2 * math.atan(t * ((1 - es) / (1 + es)) ** (_E / 2))
        if abs(nxt - phi) < 1e-12:
            phi = nxt
            break
        phi = nxt
    return lon, math.degrees(phi)


def _parse_wkb(buf: bytes, off: int = 0):
    bo = "<" if buf[off] == 1 else ">"
    (gtype,) = struct.unpack_from(bo + "I", buf, off + 1)
    off += 5
    has_z = gtype in (1002, 1005) or bool(gtype & 0x80000000)
    has_m = gtype in (2002, 2005, 3002, 3005) or bool(gtype & 0x40000000)
    base = (gtype & 0x0FFFFFFF) % 1000
    if gtype in (3002, 3005):
        has_z = has_m = True
    dim = 2 + has_z + has_m
    if base == 2:  # LineString
        (n,) = struct.unpack_from(bo + "I", buf, off)
        off += 4
        pts = [struct.unpack_from(bo + "d" * dim, buf, off + i * 8 * dim)[:2] for i in range(n)]
        return [pts], off + n * 8 * dim
    if base == 5:  # MultiLineString
        (n,) = struct.unpack_from(bo + "I", buf, off)
        off += 4
        lines = []
        for _ in range(n):
            ls, off = _parse_wkb(buf, off)
            lines.extend(ls)
        return lines, off
    raise ValueError(f"unsupported WKB type {gtype}")


def gpkg_lines(blob: bytes) -> list[list[tuple[float, float]]]:
    if blob[:2] != b"GP":
        raise ValueError("not a GeoPackage geometry blob")
    flags = blob[3]
    env = {0: 0, 1: 32, 2: 48, 3: 48, 4: 64}[(flags >> 1) & 0b111]
    lines, _ = _parse_wkb(blob, 8 + env)
    return lines


def _hav(lon1, lat1, lon2, lat2):
    r = 6371.0088
    p1, p2 = math.radians(lat1), math.radians(lat2)
    a = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(lon2 - lon1) / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def itinere_segments(bboxes: dict[str, tuple] | None = None) -> pd.DataFrame:
    """One row per segment with geodesic length and region tag (first matching bbox)."""
    con = sqlite3.connect(MANUAL["itinere"])
    q = ("SELECT fid, geom, Name, Type, Segment_s, Lower_Date, Upper_Date, Low_Date_E, Up_Date_E, "
         "Citation, Bibliograp FROM itinere_roads_simplify")
    rows = []
    for fid, geom, name, typ, cert, lo, hi, loe, hie, cit, bib in con.execute(q):
        lines = [[mercator3395_to_lonlat(x, y) for x, y in ln] for ln in gpkg_lines(geom)]
        km = sum(_hav(*ln[i], *ln[i + 1]) for ln in lines for i in range(len(ln) - 1))
        region = "rest"
        if bboxes:
            for rname, (x0, y0, x1, y1) in bboxes.items():
                if any(x0 <= lon <= x1 and y0 <= lat <= y1 for ln in lines for lon, lat in ln):
                    region = rname
                    break
        rows.append(dict(fid=fid, name=name, type=typ, certainty=cert,
                         lower_date=None if lo == 9999 else lo, upper_date=None if hi == 9999 else hi,
                         km=km, region=region, bibliography=bib))
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Harper 2016 prices
# ---------------------------------------------------------------------------
def harper_wheat_1_3c() -> pd.DataFrame:
    wb = openpyxl.load_workbook(MANUAL["harper_prices"], read_only=True, data_only=True)
    rows = list(wb["2"].iter_rows(values_only=True))
    hdr = [h for h in rows[0] if h is not None]
    df = pd.DataFrame([r[:len(hdr)] for r in rows[1:] if any(c is not None for c in r[:4])], columns=hdr)
    df = df.rename(columns={"Nominal_Price (drachmas_per_artabas)": "price_dr_art"})
    for c in ("Date_Earliest", "Date_Latest", "Year_value", "price_dr_art",
              "Drachmas_per_artabas_min", "Drachmas_per_artabas_max"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df["date_span"] = df.Date_Latest - df.Date_Earliest
    df["harper_sheet"] = "2"
    return df


def break_scan(years: np.ndarray, log_p: np.ndarray, candidates: range) -> pd.DataFrame:
    """Two-regime step model log p = a + b*1[year > tau]; SSE for each candidate tau."""
    out = []
    for tau in candidates:
        post = years > tau
        if post.sum() < 3 or (~post).sum() < 3:
            continue
        a, b = log_p[~post].mean(), log_p[post].mean()
        sse = ((log_p[~post] - a) ** 2).sum() + ((log_p[post] - b) ** 2).sum()
        out.append(dict(tau=tau, sse=sse, ratio=math.exp(b - a), n_pre=int((~post).sum()), n_post=int(post.sum())))
    return pd.DataFrame(out)


def harper_break_analysis(seed: int, n_boot: int = 2000, max_span: float | None = None,
                          window: tuple[int, int] = (1, 275), candidates: range = range(140, 221)) -> dict:
    df = harper_wheat_1_3c()
    df = df[df.Year_value.between(*window) & df.price_dr_art.notna()]
    if max_span is not None:
        df = df[df.date_span.fillna(0) <= max_span]
    y, lp = df.Year_value.to_numpy(float), np.log(df.price_dr_art.to_numpy(float))
    scan = break_scan(y, lp, candidates)
    best = scan.loc[scan.sse.idxmin()]
    # flat model SSE for an F-type comparison
    sse0 = ((lp - lp.mean()) ** 2).sum()
    rng = np.random.default_rng(seed)
    taus, ratios = [], []
    for _ in range(n_boot):
        idx = rng.integers(0, len(y), len(y))
        s = break_scan(y[idx], lp[idx], candidates)
        if len(s):
            b = s.loc[s.sse.idxmin()]
            taus.append(b.tau)
            ratios.append(b.ratio)
    pre = df[df.Year_value.between(70, 169)].price_dr_art
    post = df[df.Year_value.between(190, 259)].price_dr_art
    return dict(
        n_obs=int(len(df)), max_span=max_span,
        best_tau=int(best.tau), best_ratio=float(best.ratio), sse_break=float(best.sse), sse_flat=float(sse0),
        r2_gain=float(1 - best.sse / sse0),
        boot_tau_p05_p50_p95=[float(np.percentile(taus, q)) for q in (5, 50, 95)],
        boot_ratio_p05_p50_p95=[float(np.percentile(ratios, q)) for q in (5, 50, 95)],
        n_70_169=int(len(pre)), median_70_169=float(pre.median()) if len(pre) else None,
        share_70_169_within_6_12=float(pre.between(6, 12).mean()) if len(pre) else None,
        n_190_259=int(len(post)), median_190_259=float(post.median()) if len(post) else None,
        share_190_259_within_12_20=float(post.between(12, 20).mean()) if len(post) else None,
        seed=seed, n_boot=n_boot,
        caveat="Harper's 1st-3rd c. wheat prices overlap Rathbone's sources (same papyri): a re-derivation, not independent validation.",
    )


# ---------------------------------------------------------------------------
# Hanson 2016 cities
# ---------------------------------------------------------------------------
def hanson_tables() -> dict[str, pd.DataFrame]:
    wb = openpyxl.load_workbook(MANUAL["hanson_cities"], read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hdr = list(rows[0])
        out[ws.title] = pd.DataFrame([r for r in rows[1:] if any(c is not None for c in r)], columns=hdr)
    cities = out["Cities"]
    for c in ("Start Date", "End Date", "Longitude (X)", "Latitude (Y)"):
        cities[c] = pd.to_numeric(cities[c].replace("NULL", None), errors="coerce")
    out["Areas"]["Area"] = pd.to_numeric(out["Areas"]["Area"], errors="coerce")
    return out


def hanson_city_summary(windows: dict[str, tuple[int, int]]) -> tuple[pd.DataFrame, pd.DataFrame]:
    t = hanson_tables()
    c = t["Cities"].merge(t["Areas"], on="Primary Key", how="left")
    rows = []
    for w, (a, b) in windows.items():
        active = c[(c["Start Date"] <= b) & (c["End Date"].isna() | (c["End Date"] >= a))]
        for prov in ("ALL", "Aegyptus"):
            sub = active if prov == "ALL" else active[active.Province == prov]
            rows.append(dict(window=w, province=prov, cities=len(sub), with_area=int(sub.Area.notna().sum()),
                             area_ha_sum=float(sub.Area.sum()), area_ha_median=float(sub.Area.median())
                             if sub.Area.notna().any() else None))
    eg = c[c.Province == "Aegyptus"][["Primary Key", "Ancient Toponym", "Start Date", "End Date",
                                      "Longitude (X)", "Latitude (Y)", "Area", "Estimate Basis"]]
    return pd.DataFrame(rows), eg


def itinere_graph(bbox: tuple, snap_deg: float = 0.002, period: tuple[int, int] | None = None):
    """Undirected road graph of Itiner-e segments touching ``bbox``.

    Endpoints are snapped to a ``snap_deg`` grid (~200 m) so that segments that
    meet in the source but differ by digitising noise connect. If ``period`` is
    given, segments whose explicit date range does not overlap it are dropped
    (undated segments are kept and flagged in the edge data).
    """
    import networkx as nx
    con = sqlite3.connect(MANUAL["itinere"])
    g = nx.Graph()
    for fid, geom, name, cert, lo, hi in con.execute(
            "SELECT fid, geom, Name, Segment_s, Lower_Date, Upper_Date FROM itinere_roads_simplify"):
        lo = None if lo == 9999 else lo
        hi = None if hi == 9999 else hi
        if period and lo is not None and hi is not None and not (lo <= period[1] and period[0] <= hi):
            continue
        for ln in gpkg_lines(geom):
            pts = [mercator3395_to_lonlat(x, y) for x, y in ln]
            if not any(bbox[0] <= a <= bbox[2] and bbox[1] <= b <= bbox[3] for a, b in pts):
                continue
            km = sum(_hav(*pts[i], *pts[i + 1]) for i in range(len(pts) - 1))
            u = (round(pts[0][0] / snap_deg) * snap_deg, round(pts[0][1] / snap_deg) * snap_deg)
            v = (round(pts[-1][0] / snap_deg) * snap_deg, round(pts[-1][1] / snap_deg) * snap_deg)
            if u == v:
                continue
            if not g.has_edge(u, v) or g[u][v]["km"] > km:
                g.add_edge(u, v, km=km, fid=fid, name=name, certainty=cert, dated=lo is not None)
    return g


def nearest_node(g, lon: float, lat: float):
    return min(g.nodes, key=lambda n: _hav(lon, lat, n[0], n[1]))


def route_km(g, a: tuple[float, float], b: tuple[float, float]) -> dict:
    import networkx as nx
    na, nb = nearest_node(g, *a), nearest_node(g, *b)
    try:
        path = nx.shortest_path(g, na, nb, weight="km")
    except nx.NetworkXNoPath:
        return dict(km=None, connected=False)
    edges = [g[path[i]][path[i + 1]] for i in range(len(path) - 1)]
    km = sum(e["km"] for e in edges)
    cert = {}
    for e in edges:
        cert[e["certainty"]] = cert.get(e["certainty"], 0) + e["km"]
    return dict(km=km, connected=True, snap_gap_a_km=_hav(*a, *na), snap_gap_b_km=_hav(*b, *nb),
                certain_share=cert.get("Certain", 0) / km if km else None,
                conjectured_share=cert.get("Conjectured", 0) / km if km else None,
                dated_share=sum(e["km"] for e in edges if e["dated"]) / km if km else None,
                names=sorted({e["name"] for e in edges if e["name"]}))
